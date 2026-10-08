import logging
from pathlib import Path
import json
import csv
from io import StringIO
import pymupdf
from io import BytesIO
from docx import Document as DocxDocument

from langchain_core.documents import Document


logger = logging.getLogger(__name__)


class DocumentLoader:

    def load(
        self,
        blob_name: str,
        blob_data: bytes
    ) -> list[Document]:

        extension = Path(blob_name).suffix.lower()

        if extension == ".txt":
            return self._load_text(blob_name, blob_data)

        if extension == ".md":
            return self._load_markdown(blob_name, blob_data)

        if extension == ".json":
            return self._load_json(blob_name, blob_data)

        if extension == ".csv":
            return self._load_csv(blob_name, blob_data)
        
        if extension == ".pdf":
            return self._load_pdf(blob_name, blob_data)

        if extension == ".docx":
            return self._load_docx(blob_name, blob_data)

        raise ValueError(
            f"Unsupported file type: {extension}"
        )


    def _load_text(
        self,
        blob_name: str,
        blob_data: bytes
    ) -> list[Document]:

        text = blob_data.decode("utf-8")

        document = Document(
            page_content=text,
            metadata={
                "source": blob_name,
                "document_type": "incident_report"
            }
        )

        logger.info(
            "Loaded TXT document: %s",
            blob_name
        )

        return [document]


    def _load_markdown(
        self,
        blob_name: str,
        blob_data: bytes
    ) -> list[Document]:

        text = blob_data.decode("utf-8")

        document = Document(
            page_content=text,
            metadata={
                "source": blob_name,
                "document_type": "release_notes"
            }
        )

        logger.info(
            "Loaded Markdown document: %s",
            blob_name
        )

        return [document]

    def _load_json(
        self,
        blob_name: str,
        blob_data: bytes
        ) -> list[Document]:

        data = json.loads(
            blob_data.decode("utf-8")
        )

        documents = []
       
        for issue in data.get("known_issues", []):

            symptoms = issue.get("symptoms", [])
            
            symptoms_text = "; ".join(symptoms)
            
            page_content = (
                f"Error Code: {issue.get('error_code', '')}\n"
                f"Title: {issue.get('title', '')}\n"
                f"Symptoms: {symptoms_text}\n"
                f"Resolution: {issue.get('resolution', '')}"
            )

            document = Document(
                page_content=page_content,
                metadata={
                    "source": blob_name,
                    "document_type": "known_issue",
                    "error_code": issue.get("error_code"),
                    "topic": issue.get("topic"),
                }
            )

            documents.append(document)

        logger.info(
            "Loaded %d known issues from: %s",
            len(documents),
            blob_name
        )

        return documents

    def _load_csv(
    self,
    blob_name: str,
    blob_data: bytes
) -> list[Document]:

        text = blob_data.decode("utf-8")

        reader = csv.DictReader(
            StringIO(text)
        )

        documents = []

        for row in reader:

            fields = [
                    ("Ticket ID", row.get("ticket_id")),
                    ("Customer Issue", row.get("customer_issue")),
                    ("Error Code", row.get("error_code")),
                    ("Product Version", row.get("product_version")),
                    ("Root Cause", row.get("root_cause")),
                    ("Resolution", row.get("resolution")),
                ]

            page_content = "\n".join(
                f"{label}: {value.strip()}"
                for label, value in fields
                if value and value.strip()
            )

            document = Document(
                page_content=page_content,
                metadata={
                    "source": blob_name,
                    "document_type": "support_ticket",
                    "ticket_id": row.get("ticket_id"),
                    "topic": row.get("topic"),
                    "error_code": row.get("error_code"),
                    "status": row.get("status"),
                    "product_version": row.get("product_version"),
                }
            )

            documents.append(document)

        logger.info(
            "Loaded %d support tickets from: %s",
            len(documents),
            blob_name
        )

        return documents

    def _load_pdf(
    self,
    blob_name: str,
    blob_data: bytes
) -> list[Document]:

        documents = []

        with pymupdf.open(
            stream=blob_data,
            filetype="pdf"
        ) as pdf:

            for page_number, page in enumerate(pdf, start=1):

                text = page.get_text()

                if not text.strip():
                    continue

                document = Document(
                    page_content=text,
                    metadata={
                        "source": blob_name,
                        "document_type": "product_documentation",
                        "page": page_number,
                    }
                )

                documents.append(document)

        logger.info(
            "Loaded %d PDF pages from: %s",
            len(documents),
            blob_name
        )

        return documents

    def _load_docx(
    self,
    blob_name: str,
    blob_data: bytes
) -> list[Document]:

        docx_file = DocxDocument(
            BytesIO(blob_data)
        )

        paragraphs = [
            paragraph.text
            for paragraph in docx_file.paragraphs
            if paragraph.text.strip()
        ]

        text = "\n".join(paragraphs)

        document = Document(
            page_content=text,
            metadata={
                "source": blob_name,
                "document_type": "troubleshooting_manual",
            }
        )

        logger.info(
            "Loaded DOCX document: %s",
            blob_name
        )

        return [document]

document_loader = DocumentLoader()